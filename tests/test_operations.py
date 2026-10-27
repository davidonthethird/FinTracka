from __future__ import annotations

import plistlib
import json
import sqlite3
import subprocess
import stat
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fintracka.backups import backup_database, restore_database, verify_backup
from fintracka.db import connect, migrate
from fintracka.exports import SHEET_ORDER, build_snapshot, export_excel, export_google
from fintracka.google_sheets import (
    DRIVE_FILE_SCOPE,
    GoogleSheetsGateway,
    authorize_google,
    configure_google,
    google_status,
)
from fintracka.keychain import KeychainSecrets
from fintracka.launchd import (
    build_launch_agent,
    install_schedule,
    remove_schedule,
    schedule_status,
)
from fintracka.locks import ProcessLock, ProcessLockError
from fintracka.manual_cash import create_manual_cash
from fintracka.planning import create_income_source, set_recurring_target
from fintracka.providers import NormalizedTransaction, SyncBatch
from fintracka.savings import create_manual_asset, deactivate_manual_asset
from fintracka.subcategories import create_subcategory
from fintracka.sync import apply_batch
from fintracka.synthetic import create_synthetic_ledger, synthetic_batches


@pytest.fixture
def ledger_path(tmp_path: Path) -> Path:
    path = tmp_path / "ledger.sqlite3"
    connection = connect(path)
    migrate(connection)
    connection_id = create_synthetic_ledger(connection, count=1)[0]
    apply_batch(connection, connection_id, synthetic_batches([connection_id])[connection_id])
    connection.close()
    return path


def test_backup_is_private_verified_and_restores_into_empty_target(
    ledger_path: Path, tmp_path: Path
) -> None:
    backup = backup_database(ledger_path, tmp_path / "backups", keep=30)
    assert backup.stat().st_mode & 0o777 == 0o600

    source = connect(ledger_path)
    source_count = source.execute("SELECT COUNT(*) FROM transactions").fetchone()[0]
    source.close()
    restored = tmp_path / "restored" / "ledger.sqlite3"
    restore_database(backup, restored, verify=True)
    connection = connect(restored)
    assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    assert connection.execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == source_count
    connection.close()


def test_pre_migration_backup_is_verified_and_private(tmp_path: Path) -> None:
    path = tmp_path / "ledger.sqlite3"
    connection = connect(path)

    migrate(connection)

    backups = sorted((tmp_path / "backups").glob("fintracka-pre-migration-*.sqlite3"))
    assert len(backups) == 6
    for backup in backups:
        verify_backup(backup)
        assert stat.S_IMODE(backup.stat().st_mode) == 0o600
    assert connection.execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0] == 18
    connection.close()


def test_pre_migration_backup_failure_leaves_version_12_unchanged(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "ledger.sqlite3"
    connection = connect(path)

    def fail_backup(*args, **kwargs):  # type: ignore[no-untyped-def]
        raise sqlite3.DatabaseError("verification failed")

    monkeypatch.setattr("fintracka.db.backup_connection", fail_backup, raising=False)
    with pytest.raises(sqlite3.DatabaseError, match="verification failed"):
        migrate(connection)

    assert connection.execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0] == 12
    assert "credential_profile" not in {
        row[1] for row in connection.execute("PRAGMA table_info(provider_connections)")
    }
    connection.close()


def test_excel_export_is_deterministic_replaces_deleted_rows_and_escapes_formulas(
    ledger_path: Path, tmp_path: Path
) -> None:
    connection = connect(ledger_path)
    connection_id = connection.execute(
        "SELECT id FROM provider_connections WHERE provider = 'synthetic'"
    ).fetchone()[0]
    dangerous = NormalizedTransaction(
        "formula", "checking-1", "2026-10-10T12:00:00Z", -100, "USD", "=HYPERLINK(\"https://evil\")"
    )
    apply_batch(connection, connection_id, SyncBatch(added=(dangerous,), next_checkpoint="export"))
    formula_id = connection.execute(
        "SELECT transaction_id FROM transaction_external_ids WHERE external_id = 'formula'"
    ).fetchone()[0]
    output = tmp_path / "FinTracka.xlsx"
    export_excel(connection, output)

    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == list(SHEET_ORDER)
    assert workbook["_FinTrackaMetadata"].sheet_state == "hidden"
    transaction_rows = list(workbook["Transactions"].values)
    assert "Spending owner" in transaction_rows[0]
    assert len({row[0] for row in transaction_rows[1:]}) == len(transaction_rows) - 1
    formula_row = next(row for row in transaction_rows[1:] if row[0] == formula_id)
    assert formula_row[5].startswith("'=")
    first_rows = transaction_rows
    workbook.close()

    connection.execute("UPDATE transactions SET tombstoned = 1 WHERE id = ?", (formula_id,))
    export_excel(connection, output)
    workbook = load_workbook(output, data_only=False)
    second_rows = list(workbook["Transactions"].values)
    assert formula_id not in {row[0] for row in second_rows[1:]}
    assert second_rows == [row for row in first_rows if row[0] != formula_id]
    workbook.close()
    connection.close()


def test_google_export_replaces_complete_sheets_in_bounded_batches(ledger_path: Path) -> None:
    connection = connect(ledger_path)
    gateway = FakeSheetsGateway()
    export_google(connection, gateway, batch_size=2)
    first = {name: rows[:] for name, rows in gateway.sheets.items() if name != "_FinTrackaMetadata"}
    export_google(connection, gateway, batch_size=2)

    assert tuple(gateway.sheets) == SHEET_ORDER
    assert {name: rows for name, rows in gateway.sheets.items() if name != "_FinTrackaMetadata"} == first
    assert all(size <= 2 for size in gateway.batch_sizes)
    assert gateway.protected == set(SHEET_ORDER)
    connection.close()


def test_export_includes_effective_planning_and_manual_cash_deterministically(
    ledger_path: Path,
) -> None:
    connection = connect(ledger_path)
    salary = create_income_source(connection, "Paycheck", "Salary", 590_001, "2026-01")
    custom = create_income_source(connection, "=Rental", "Custom", 110_000, "2026-01")
    shopping = connection.execute("SELECT id FROM categories WHERE name = 'Shopping'").fetchone()[0]
    tzedaka = connection.execute("SELECT id FROM categories WHERE name = 'Tzedaka'").fetchone()[0]
    set_recurring_target(connection, shopping, 40_000, "2026-01")
    connection.execute("UPDATE categories SET name = '=Shopping formula' WHERE id = ?", (shopping,))
    manual_id = create_manual_cash(connection, "Cash gift", 2_500, tzedaka, "2026-10-13", "Shared")

    first = build_snapshot(connection)
    second = build_snapshot(connection)

    assert "Planned Income" in SHEET_ORDER
    assert first["Planned Income"] == second["Planned Income"]
    income_rows = first["Planned Income"]
    assert {row[0] for row in income_rows[1:]} >= {f"{salary}:2026-01", f"{custom}:2026-12"}
    rental = next(row for row in income_rows[1:] if row[0] == f"{custom}:2026-01")
    assert rental[2] == "'=Rental" and rental[3] == "Custom"

    monthly = first["Monthly Budget"]
    assert len({row[1] for row in monthly[1:]}) == 12
    october_shopping = next(
        row for row in monthly[1:] if row[1] == "2026-10" and row[2] == "'=Shopping formula"
    )
    assert october_shopping[3:12] == [0, 40_000, 0, 40_000, 0, 700_001, 590_001, 12, "USD"]
    october_tzedaka = next(
        row for row in monthly[1:] if row[1] == "2026-10" and row[2] == "Tzedaka"
    )
    assert october_tzedaka[4] == 59_001
    assert october_tzedaka[5] == 2_500
    yearly_tzedaka = next(row for row in first["Yearly Budget"][1:] if row[2] == "Tzedaka")
    assert yearly_tzedaka[4] == 59_001 * 12 and yearly_tzedaka[5] == 2_500

    transaction_rows = first["Transactions"]
    manual = next(row for row in transaction_rows[1:] if row[0] == manual_id)
    assert transaction_rows[0][-1] == "Manual cash"
    assert manual[-1] == "Yes"
    assert [row[0] for row in transaction_rows[1:]].count(manual_id) == 1
    connection.close()


def test_export_includes_stable_subcategories_and_active_manual_assets(
    ledger_path: Path,
) -> None:
    connection = connect(ledger_path)
    transaction = connection.execute("SELECT id FROM transactions LIMIT 1").fetchone()
    assert transaction is not None
    category_id = connection.execute(
        "SELECT id FROM categories WHERE category_type = 'expense' ORDER BY sort_order LIMIT 1"
    ).fetchone()[0]
    connection.execute(
        "UPDATE transactions SET category_id = ?, category_source = 'manual' WHERE id = ?",
        (category_id, transaction["id"]),
    )
    subcategory_id = create_subcategory(connection, category_id, "+Home detail")
    connection.execute(
        "UPDATE categories SET name = '@Parent category' WHERE id = ?", (category_id,)
    )
    connection.execute(
        "UPDATE transactions SET subcategory_id = ?, subcategory_source = 'manual' WHERE id = ?",
        (subcategory_id, transaction["id"]),
    )
    active_asset = create_manual_asset(
        connection, "=Emergency reserve", "savings", "Shared", 123_456, "2026-10-14"
    )
    inactive_asset = create_manual_asset(
        connection, "+Old reserve", "other", "David", 200, "2026-10-01"
    )
    inactive_updated = connection.execute(
        "SELECT updated_at FROM manual_assets WHERE id = ?", (inactive_asset,)
    ).fetchone()[0]
    deactivate_manual_asset(connection, inactive_asset, inactive_updated)

    first = build_snapshot(connection)
    second = build_snapshot(connection)

    transaction_header = first["Transactions"][0]
    assert transaction_header[7:11] == [
        "Category ID", "Category", "Subcategory ID", "Subcategory"
    ]
    exported_transaction = next(
        row for row in first["Transactions"][1:] if row[0] == transaction["id"]
    )
    assert exported_transaction[7:11] == [
        category_id, "'@Parent category", subcategory_id, "'+Home detail"
    ]
    assert first["Subcategories"] == second["Subcategories"]
    assert first["Manual Assets"] == second["Manual Assets"]
    assert [row[0] for row in first["Manual Assets"][1:]] == [active_asset]
    assert first["Manual Assets"][1][1] == "'=Emergency reserve"

    active_updated = connection.execute(
        "SELECT updated_at FROM manual_assets WHERE id = ?", (active_asset,)
    ).fetchone()[0]
    deactivate_manual_asset(connection, active_asset, active_updated)
    assert build_snapshot(connection)["Manual Assets"] == [first["Manual Assets"][0]]
    connection.close()


def test_google_authorization_keeps_credentials_in_keychain_and_workbook_id_in_sqlite(
    ledger_path: Path,
) -> None:
    connection = connect(ledger_path)
    secrets = KeychainSecrets(FakeKeyring())
    client_config = {
        "installed": {
            "client_id": "desktop-client.apps.googleusercontent.com",
            "client_secret": "local-client-secret",
            "auth_uri": "https://accounts.google.com/o/oauth2/auth",
            "token_uri": "https://oauth2.googleapis.com/token",
            "redirect_uris": ["http://localhost"],
        }
    }
    configure_google(connection, secrets, json.dumps(client_config))
    gateway = AuthorizingGateway()
    authorize_google(
        connection,
        secrets,
        oauth_authorizer=lambda config, scopes: (
            json.dumps({"refresh_token": "refresh-secret", "client_id": config["installed"]["client_id"]})
            if scopes == [DRIVE_FILE_SCOPE]
            else ""
        ),
        gateway_factory=lambda credentials: gateway,
    )

    status = google_status(connection)
    assert status == {"configured": True, "authorized": True, "workbook_id": "workbook-123"}
    assert gateway.created == [list(SHEET_ORDER)]
    database_dump = "\n".join(
        str(value)
        for row in connection.iterdump()
        for value in [row]
    )
    assert "local-client-secret" not in database_dump
    assert "refresh-secret" not in database_dump
    assert secrets.get("google:oauth:credentials") is not None
    connection.close()


def test_google_gateway_clears_stale_rows_batches_updates_and_replaces_protection() -> None:
    service = FakeSheetsService()
    gateway = GoogleSheetsGateway(service, "workbook-123")
    rows = [["ID", "Value"], ["1", "first"], ["2", "second"]]

    gateway.replace_sheet("Transactions", rows, batch_size=2)
    gateway.protect_sheet("Transactions", "overwritten on sync")

    assert service.clears == [("workbook-123", "'Transactions'")]
    assert service.updates == [
        ("workbook-123", "'Transactions'!A1", rows[:2]),
        ("workbook-123", "'Transactions'!A3", rows[2:]),
    ]
    requests = service.batch_requests[-1]
    assert requests[0] == {"deleteProtectedRange": {"protectedRangeId": 7}}
    assert requests[1]["addProtectedRange"]["protectedRange"]["warningOnly"] is True


def test_process_lock_is_nonblocking_and_launchd_schedule_is_exact(tmp_path: Path) -> None:
    lock_path = tmp_path / "fintracka.lock"
    with ProcessLock(lock_path):
        with pytest.raises(ProcessLockError):
            with ProcessLock(lock_path):
                pass

    payload = plistlib.loads(
        build_launch_agent(Path("/usr/bin/python3"), Path("/tmp/fintracka-data"))
    )
    assert payload["Label"] == "com.fintracka.sync"
    assert payload["StartCalendarInterval"] == {"Hour": 6, "Minute": 17}
    assert payload["RunAtLoad"] is True
    assert payload["ProgramArguments"][-1] == "/tmp/fintracka-data"


def test_launchd_schedule_install_status_and_remove_are_explicit(tmp_path: Path) -> None:
    destination = tmp_path / "Library" / "LaunchAgents" / "com.fintracka.sync.plist"
    calls: list[list[str]] = []

    def runner(command: list[str], **kwargs):  # type: ignore[no-untyped-def]
        del kwargs
        calls.append(command)
        return subprocess.CompletedProcess(command, 0, "", "")

    installed = install_schedule(
        destination,
        Path("/opt/fintracka/bin/python"),
        tmp_path / "data",
        runner=runner,
        user_id=501,
    )
    assert installed == destination and destination.stat().st_mode & 0o777 == 0o600
    assert calls[-1][:3] == ["launchctl", "bootstrap", "gui/501"]
    assert schedule_status(destination, runner=runner, user_id=501) == {
        "installed": True,
        "loaded": True,
    }

    remove_schedule(destination, runner=runner, user_id=501)
    assert calls[-1][:3] == ["launchctl", "bootout", "gui/501"]
    assert not destination.exists()


class FakeSheetsGateway:
    def __init__(self) -> None:
        self.sheets: dict[str, list[list[object]]] = {}
        self.batch_sizes: list[int] = []
        self.protected: set[str] = set()

    def replace_sheet(self, name: str, rows: list[list[object]], batch_size: int) -> None:
        self.batch_sizes.extend(
            len(rows[index : index + batch_size]) for index in range(0, len(rows), batch_size)
        )
        self.sheets[name] = rows

    def protect_sheet(self, name: str, warning: str) -> None:
        assert "overwritten" in warning
        self.protected.add(name)


class FakeKeyring:
    def __init__(self) -> None:
        self.values: dict[tuple[str, str], str] = {}

    def set_password(self, service: str, account: str, password: str) -> None:
        self.values[(service, account)] = password

    def get_password(self, service: str, account: str) -> str | None:
        return self.values.get((service, account))

    def delete_password(self, service: str, account: str) -> None:
        self.values.pop((service, account), None)


class AuthorizingGateway:
    def __init__(self) -> None:
        self.created: list[list[str]] = []

    def create_workbook(self, names: tuple[str, ...]) -> str:
        self.created.append(list(names))
        return "workbook-123"


class FakeExecutable:
    def __init__(self, value: dict[str, object], callback=None) -> None:  # type: ignore[no-untyped-def]
        self.value = value
        self.callback = callback

    def execute(self) -> dict[str, object]:
        if self.callback:
            self.callback()
        return self.value


class FakeValuesResource:
    def __init__(self, owner: "FakeSheetsService") -> None:
        self.owner = owner

    def clear(self, *, spreadsheetId: str, range: str, body: dict[str, object]):  # type: ignore[no-untyped-def]
        assert body == {}
        return FakeExecutable({}, lambda: self.owner.clears.append((spreadsheetId, range)))

    def update(self, *, spreadsheetId: str, range: str, valueInputOption: str, body: dict[str, object]):  # type: ignore[no-untyped-def]
        assert valueInputOption == "RAW"
        return FakeExecutable(
            {},
            lambda: self.owner.updates.append((spreadsheetId, range, body["values"])),
        )


class FakeSpreadsheetsResource:
    def __init__(self, owner: "FakeSheetsService") -> None:
        self.owner = owner
        self.values_resource = FakeValuesResource(owner)

    def values(self) -> FakeValuesResource:
        return self.values_resource

    def get(self, *, spreadsheetId: str, fields: str):  # type: ignore[no-untyped-def]
        assert spreadsheetId == "workbook-123"
        assert "protectedRanges" in fields
        return FakeExecutable(
            {
                "sheets": [
                    {
                        "properties": {"sheetId": 10, "title": "Transactions"},
                        "protectedRanges": [{"protectedRangeId": 7, "description": "FinTracka managed range"}],
                    }
                ]
            }
        )

    def batchUpdate(self, *, spreadsheetId: str, body: dict[str, object]):  # type: ignore[no-untyped-def]
        return FakeExecutable(
            {}, lambda: self.owner.batch_requests.append(body["requests"])
        )


class FakeSheetsService:
    def __init__(self) -> None:
        self.clears: list[tuple[str, str]] = []
        self.updates: list[tuple[str, str, object]] = []
        self.batch_requests: list[list[dict[str, object]]] = []
        self.resource = FakeSpreadsheetsResource(self)

    def spreadsheets(self) -> FakeSpreadsheetsResource:
        return self.resource
