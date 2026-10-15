from __future__ import annotations

import os
import sqlite3
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Protocol

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from fintracka.ids import new_id
from fintracka.planning import planning_snapshot
from fintracka.reports import spending_by_category

SHEET_ORDER = (
    "Transactions",
    "Accounts",
    "Balances",
    "Balance History",
    "Categories",
    "Subcategories",
    "Auto-categorization Rules",
    "Manual Assets",
    "Planned Income",
    "Monthly Budget",
    "Yearly Budget",
    "Spending Trends",
    "_FinTrackaMetadata",
)


class SheetsGateway(Protocol):
    def replace_sheet(self, name: str, rows: list[list[object]], batch_size: int) -> None: ...
    def protect_sheet(self, name: str, warning: str) -> None: ...


def export_excel(connection: sqlite3.Connection, destination: str | Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    destination.parent.chmod(0o700)
    snapshot = _consistent_snapshot(connection)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEET_ORDER:
        sheet = workbook.create_sheet(name)
        for row in snapshot[name]:
            sheet.append(row)
        _format_sheet(sheet)
        if name == "_FinTrackaMetadata":
            sheet.sheet_state = "hidden"
    descriptor, temporary_name = tempfile.mkstemp(prefix=".fintracka-export-", suffix=".xlsx", dir=destination.parent)
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        workbook.save(temporary)
        workbook.close()
        validation = load_workbook(temporary, read_only=True, data_only=False)
        try:
            if tuple(validation.sheetnames) != SHEET_ORDER:
                raise ValueError("export validation failed: sheet set changed")
        finally:
            validation.close()
        temporary.chmod(0o600)
        os.replace(temporary, destination)
    finally:
        if temporary.exists():
            temporary.unlink()
    _audit(connection, "excel_export", str(destination))
    return destination


def export_google(
    connection: sqlite3.Connection,
    gateway: SheetsGateway,
    *,
    batch_size: int = 500,
) -> None:
    if batch_size < 1:
        raise ValueError("batch size must be positive")
    snapshot = _consistent_snapshot(connection)
    warning = "Edited locally in FinTracka; this sheet is overwritten on sync."
    for name in SHEET_ORDER:
        gateway.replace_sheet(name, snapshot[name], batch_size)
        gateway.protect_sheet(name, warning)
    _audit(connection, "google_sheets_export", "app_owned_workbook")


def _consistent_snapshot(connection: sqlite3.Connection) -> dict[str, list[list[object]]]:
    try:
        connection.execute("BEGIN IMMEDIATE")
        snapshot = build_snapshot(connection)
        connection.commit()
        return snapshot
    except Exception:
        connection.rollback()
        raise


def build_snapshot(connection: sqlite3.Connection) -> dict[str, list[list[object]]]:
    database_id = connection.execute("SELECT value FROM app_metadata WHERE key = 'database_instance_id'").fetchone()
    if database_id is None:
        value = new_id()
        connection.execute("INSERT INTO app_metadata (key, value) VALUES ('database_instance_id', ?)", (value,))
        database_id = (value,)
    sheets: dict[str, list[list[object]]] = {}
    sheets["Transactions"] = _rows(
        ["FinTracka ID", "Date", "Account", "Amount (minor units)", "Currency", "Description", "Merchant", "Category ID", "Category", "Subcategory ID", "Subcategory", "Notes", "Pending", "Last updated", "Spending owner", "Manual cash"],
        connection.execute(
            """
            SELECT t.id, t.local_date, a.name, t.amount_minor, t.currency,
                   t.provider_description, COALESCE(t.user_merchant, t.provider_merchant),
                   t.category_id, c.name, t.subcategory_id, s.name,
                   t.notes, t.pending, t.updated_at,
                   CASE WHEN c.category_type = 'expense' THEN COALESCE(t.spending_owner, a.spending_owner) END,
                   CASE WHEN pc.provider = 'manual' THEN 'Yes' ELSE 'No' END
            FROM transactions t JOIN accounts a ON a.id = t.account_id
            JOIN provider_connections pc ON pc.id = a.connection_id
            LEFT JOIN categories c ON c.id = t.category_id
            LEFT JOIN subcategories s ON s.id = t.subcategory_id
            WHERE t.tombstoned = 0 AND a.include_export = 1 AND a.hidden = 0
            ORDER BY t.local_date, t.id
            """
        ),
    )
    sheets["Accounts"] = _rows(
        ["FinTracka ID", "Name", "Nickname", "Type", "Subtype", "Nature", "Currency", "Status", "Spending owner", "Last updated"],
        connection.execute(
            "SELECT id, name, nickname, type, subtype, nature, currency, status, spending_owner, updated_at FROM accounts WHERE include_export = 1 AND hidden = 0 ORDER BY nature, name, id"
        ),
    )
    sheets["Balances"] = _rows(
        ["FinTracka ID", "Account", "Observed at", "Current (minor units)", "Available (minor units)", "Currency", "Source"],
        connection.execute(
            """
            SELECT b.id, a.name, b.observed_at, b.current_minor, b.available_minor, b.currency, b.source
            FROM account_balance_snapshots b JOIN accounts a ON a.id = b.account_id
            WHERE a.include_export = 1 AND a.hidden = 0 AND b.observed_at = (
                SELECT MAX(latest.observed_at) FROM account_balance_snapshots latest WHERE latest.account_id = a.id
            ) ORDER BY a.name, b.id
            """
        ),
    )
    sheets["Balance History"] = _rows(
        ["FinTracka ID", "Account", "Observed at", "Current (minor units)", "Available (minor units)", "Currency", "Source"],
        connection.execute(
            "SELECT b.id, a.name, b.observed_at, b.current_minor, b.available_minor, b.currency, b.source FROM account_balance_snapshots b JOIN accounts a ON a.id = b.account_id WHERE a.include_export = 1 AND a.hidden = 0 ORDER BY b.observed_at, a.name, b.id"
        ),
    )
    sheets["Categories"] = _rows(
        ["FinTracka ID", "Group", "Category", "Type", "Reporting behavior", "Archived", "Last updated"],
        connection.execute(
            "SELECT c.id, g.name, c.name, c.category_type, c.reporting_behavior, c.archived, c.updated_at FROM categories c LEFT JOIN category_groups g ON g.id = c.group_id ORDER BY c.sort_order, c.name, c.id"
        ),
    )
    sheets["Subcategories"] = _rows(
        ["FinTracka ID", "Category ID", "Category", "Subcategory", "Sort order", "Archived", "Last updated"],
        connection.execute(
            """
            SELECT s.id, s.category_id, c.name, s.name, s.sort_order, s.archived, s.updated_at
            FROM subcategories s JOIN categories c ON c.id = s.category_id
            ORDER BY c.sort_order, c.name, s.sort_order, s.name, s.id
            """
        ),
    )
    sheets["Auto-categorization Rules"] = _rows(
        ["FinTracka ID", "Priority", "Name", "Enabled", "Conditions", "Category", "Last updated"],
        connection.execute(
            "SELECT r.id, r.priority, r.name, r.enabled, r.conditions_json, c.name, r.updated_at FROM categorization_rules r JOIN categories c ON c.id = r.category_id ORDER BY r.priority, r.id"
        ),
    )
    sheets["Manual Assets"] = _rows(
        ["FinTracka ID", "Title", "Kind", "Owner", "Current value (minor units)", "Currency", "As of date", "Last updated"],
        connection.execute(
            """
            SELECT id, title, kind, owner, current_value_minor, currency, as_of_date, updated_at
            FROM manual_assets WHERE active = 1
            ORDER BY kind, title, id
            """
        ),
    )
    export_year = datetime.now().year
    planned_income, monthly_budget, yearly_budget = _planning_rows(connection, export_year)
    sheets["Planned Income"] = planned_income
    sheets["Monthly Budget"] = monthly_budget
    sheets["Yearly Budget"] = yearly_budget
    trends: list[list[object]] = [["FinTracka ID", "Month", "Category", "Spending (minor units)", "Last updated"]]
    months = [row[0] for row in connection.execute("SELECT DISTINCT substr(local_date, 1, 7) FROM transactions WHERE tombstoned = 0 ORDER BY 1")]
    exported_at = _now()
    for month in months:
        for row in spending_by_category(connection, month):
            trends.append([f"{row['category_id']}:{month}", month, row["category"], row["amount_minor"], exported_at])
    sheets["Spending Trends"] = trends
    row_counts = {name: max(len(rows) - 1, 0) for name, rows in sheets.items()}
    sheets["_FinTrackaMetadata"] = [
        ["Key", "Value"],
        ["schema_version", connection.execute("SELECT MAX(version) FROM schema_migrations").fetchone()[0]],
        ["database_instance_id", database_id[0]],
        ["exported_at", exported_at],
        *[[f"rows:{name}", count] for name, count in row_counts.items()],
    ]
    return sheets


def _planning_rows(
    connection: sqlite3.Connection, year: int
) -> tuple[list[list[object]], list[list[object]], list[list[object]]]:
    income_rows: list[list[object]] = [[
        "FinTracka ID", "Month", "Title", "Type", "Amount (minor units)",
        "Currency", "Last updated",
    ]]
    monthly_rows: list[list[object]] = [[
        "FinTracka ID", "Month", "Category", "Historical average (minor units)",
        "Target (minor units)", "Actual (minor units)", "Remaining (minor units)",
        "Tzedaka minimum (minor units)", "Planned income (minor units)",
        "Salary (minor units)", "Average window (months)", "Currency", "Last updated",
    ]]
    category_updated = {
        row["id"]: row["updated_at"]
        for row in connection.execute("SELECT id, updated_at FROM categories")
    }
    monthly_snapshots = []
    for number in range(1, 13):
        month = f"{year:04d}-{number:02d}"
        snapshot = planning_snapshot(connection, month, 12)
        monthly_snapshots.append(snapshot)
        for income in snapshot.incomes:
            income_rows.append([
                f"{income.source_id}:{month}", month, _safe(income.title), income.income_type,
                income.amount_minor, "USD", income.updated_at,
            ])
        for category in snapshot.categories:
            monthly_rows.append([
                f"{category.category_id}:{month}", month, _safe(category.category),
                category.historical_average_minor, category.target_minor, category.actual_minor,
                category.remaining_minor, category.tzedaka_minimum_minor,
                snapshot.planned_income_total_minor, snapshot.salary_total_minor,
                snapshot.average_months, "USD", category_updated[category.category_id],
            ])

    yearly_rows: list[list[object]] = [[
        "FinTracka ID", "Year", "Category", "Historical average (minor units)",
        "Target (minor units)", "Actual (minor units)", "Remaining (minor units)",
        "Tzedaka minimum (minor units)", "Planned income (minor units)",
        "Salary (minor units)", "Average window (months)", "Currency", "Last updated",
    ]]
    if monthly_snapshots:
        for position, category in enumerate(monthly_snapshots[0].categories):
            categories = [snapshot.categories[position] for snapshot in monthly_snapshots]
            yearly_rows.append([
                f"{category.category_id}:{year}", str(year), _safe(category.category),
                _rounded_average(sum(row.historical_average_minor for row in categories), 12),
                sum(row.target_minor for row in categories),
                sum(row.actual_minor for row in categories),
                sum(row.remaining_minor for row in categories),
                sum(row.tzedaka_minimum_minor for row in categories),
                sum(snapshot.planned_income_total_minor for snapshot in monthly_snapshots),
                sum(snapshot.salary_total_minor for snapshot in monthly_snapshots),
                12, "USD", category_updated[category.category_id],
            ])
    return income_rows, monthly_rows, yearly_rows


def _rounded_average(total: int, count: int) -> int:
    return (total + count // 2) // count if total >= 0 else -((-total + count // 2) // count)


def _rows(headers: list[str], cursor: sqlite3.Cursor) -> list[list[object]]:
    return [headers, *[[_safe(value) for value in row] for row in cursor]]


def _safe(value: object) -> object:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _format_sheet(sheet) -> None:  # type: ignore[no-untyped-def]
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="142638")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Avenir Next", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in sheet.columns:
        values = [len(str(cell.value)) if cell.value is not None else 0 for cell in column]
        sheet.column_dimensions[column[0].column_letter].width = min(max(max(values, default=8) + 2, 10), 42)
        for cell in column[1:]:
            cell.font = Font(name="Avenir Next", size=10, color="142638")
            cell.alignment = Alignment(vertical="center")
    sheet["A1"].comment = Comment(
        "Edited locally in FinTracka; this sheet is overwritten on sync.", "FinTracka"
    )
    sheet.protection.sheet = True


def _audit(connection: sqlite3.Connection, action: str, entity_id: str) -> None:
    connection.execute(
        "INSERT INTO audit_events (id, actor, action, entity_type, entity_id, occurred_at) VALUES (?, 'system', ?, 'export', ?, ?)",
        (new_id(), action, entity_id, _now()),
    )


def _now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
